import pytest
from pathlib import Path

from build_your_own_rag.models import SourceDocument, ExtractionStrategy
from build_your_own_rag.parsing.xlsx_parser import parse_xlsx

@pytest.fixture
def sample_xlsx_file(tmp_path):
    # We require openpyxl for the parser anyway, so use it to write a test file
    import openpyxl
    file_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    
    # Sheet 1
    ws1 = wb.active
    ws1.title = "Employees"
    ws1.append(["Name", "Age", "Department"])
    ws1.append(["Alice", 30, "Engineering"])
    ws1.append(["Bob", 25, "Sales"])
    ws1.append([]) # empty row
    ws1.append(["Charlie", None, "Marketing"])
    
    # Sheet 2
    ws2 = wb.create_sheet("Locations")
    ws2.append(["City", "Country"])
    ws2.append(["New York", "USA"])
    ws2.append(["London", "UK"])
    
    wb.save(file_path)
    return file_path


def test_parse_xlsx(sample_xlsx_file):
    source = SourceDocument(
        source_id="test_id_xlsx",
        source_type="local_xlsx",
        path=sample_xlsx_file,
        filename="test.xlsx",
        extension="xlsx",
        size_bytes=sample_xlsx_file.stat().st_size,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    
    parsed = parse_xlsx(source, strategy=ExtractionStrategy.FAST)
    
    assert parsed.parser_name == "xlsx_native"
    assert parsed.source == source
    assert parsed.status == "success"
    
    # Check that it prepended sheet names and headers properly
    assert "[Sheet: Employees] Name: Alice | Age: 30 | Department: Engineering" in parsed.text
    assert "[Sheet: Employees] Name: Bob | Age: 25 | Department: Sales" in parsed.text
    assert "[Sheet: Employees] Name: Charlie | Department: Marketing" in parsed.text
    assert "Age" not in parsed.text.split("Charlie")[1].split("|")[0]
    
    assert "[Sheet: Locations] City: New York | Country: USA" in parsed.text
    
    xlsx_meta = parsed.metadata.format_specific["xlsx"]
    assert xlsx_meta["sheet_count"] == 2
    assert "Employees" in xlsx_meta["sheet_names"]
    assert "Locations" in xlsx_meta["sheet_names"]
    assert xlsx_meta["total_data_rows"] == 5  # 3 in Employees, 2 in Locations


@pytest.fixture
def complex_xlsx_file():
    file_path = Path(__file__).parent.parent / "data" / "test_docs" / "complex_data.xlsx"
    if not file_path.exists():
        pytest.skip("complex_data.xlsx not generated. Run generate_dummy.py first.")
    return file_path


def test_parse_complex_xlsx(complex_xlsx_file):
    source = SourceDocument(
        source_id="test_complex_xlsx",
        source_type="local_xlsx",
        path=complex_xlsx_file,
        filename="complex_data.xlsx",
        extension="xlsx",
        size_bytes=complex_xlsx_file.stat().st_size,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    
    parsed = parse_xlsx(source, strategy=ExtractionStrategy.FAST)
    
    assert parsed.parser_name == "xlsx_native"
    assert parsed.status == "success"
    
    # Check Sales Data sheet
    # Date should be parsed. openpyxl might format it as a datetime string.
    assert "[Sheet: Sales Data] Date:" in parsed.text
    
    # Check that formulas are NOT raw formulas like =C2*0.1 but actual values, 
    # except we used data_only=True which requires Excel to evaluate the formula once.
    # Because openpyxl created the formula, it hasn't been evaluated by Excel yet,
    # so openpyxl might return None for the evaluated value. Let's just check that it
    # extracted the rest of the row accurately.
    assert "[Sheet: Sales Data] Date: 2023-01-15 00:00:00 | Region: North | Total Sales: 1000" in parsed.text
    
    # Check Employee Roster sheet
    assert "[Sheet: Employee Roster] Emp ID: E001 | Full Name: Alice Smith | Department: Engineering | Notes: On leave" in parsed.text
    
    # Empty column handling
    assert "[Sheet: Employee Roster] Emp ID: E002 | Full Name: Bob Jones | Notes: Missing dept" in parsed.text
    
    # Multiline
    assert "Top performer\nQ3" in parsed.text
    
    xlsx_meta = parsed.metadata.format_specific["xlsx"]
    assert xlsx_meta["sheet_count"] == 2
    assert xlsx_meta["total_data_rows"] == 6 # 3 sales + 3 employees
